"""Small adapters for the published MarkItDown 0.1.7 converters.

Keep native text/table conversion upstream-owned. Route embedded images through
our existing vision service without depending on unreleased image hooks.
"""
import io
import re

from markitdown import StreamInfo
from markitdown.converters import DocxConverter, PptxConverter, XlsxConverter


def escape_table_cell(value):
    """Escape literal pipes and collapse record-internal newlines for display."""
    value = re.sub(r"(?<!\\)(\\*)\|", lambda m: m.group(1) * 2 + r"\|", value)
    return value.replace("\r\n", " ").replace("\r", " ").replace("\n", " ")


def office_converter(fmt, vision):
    class Docx(DocxConverter):
        def convert(self, file_stream, stream_info, **kwargs):
            import mammoth
            from mammoth import html
            from markitdown.converter_utils.docx.pre_process import pre_process_docx

            def image_element(image):
                with image.open() as stream:
                    token = vision.writer.token(vision.analyze(
                        stream, StreamInfo(mimetype=image.content_type)))
                return [html.element("p", {}, [html.text(token)])]

            prepared = pre_process_docx(file_stream)
            try:
                result = mammoth.convert_to_html(
                    prepared, style_map=kwargs.get("style_map"), convert_image=image_element)
                return self._html_converter.convert_string(result.value, **kwargs)
            finally:
                if prepared is not file_stream:
                    prepared.close()

    class Pptx(PptxConverter):
        def _get_image_info(self, shape):
            blob, mimetype, filename = super()._get_image_info(shape)
            if blob is None:
                vision.writer.warn("幻灯片中存在无法提取的图片。")
                chunks = [("metadata", "[图片未提取]")]
            else:
                with io.BytesIO(blob) as stream:
                    chunks = vision.analyze(stream, StreamInfo(mimetype=mimetype, filename=filename))
            token = vision.writer.token(chunks)
            props = shape._element._nvXxPr.cNvPr
            props.set("descr", token + " " + props.get("descr", ""))
            return blob, mimetype, filename

        def convert(self, file_stream, stream_info, **kwargs):
            result = super().convert(file_stream, stream_info, **kwargs)
            # The native converter emits image alt text. Unwrap only markers we
            # inserted, leaving ordinary user-authored Markdown untouched.
            def unwrap(match):
                if any(token in match[1] for token in vision.writer.tokens):
                    return match[1]
                return match[0]
            result.markdown = re.sub(r"!\[([^\]\n]*)\]\([^\)\n]*\)", unwrap, result.markdown)
            return result

    class Xlsx(XlsxConverter):
        def convert(self, file_stream, stream_info, **kwargs):
            from openpyxl import load_workbook
            result = super().convert(file_stream, stream_info, **kwargs)
            file_stream.seek(0)
            book = load_workbook(file_stream)
            try:
                # The bridge has already replaced sheet names with unique
                # per-call markers, so user text cannot spoof these boundaries.
                starts = [result.markdown.find("## " + name + "\n") for name in book.sheetnames]
                if any(pos < 0 for pos in starts) or starts != sorted(starts):
                    raise ValueError("Native worksheet boundaries changed")
                additions = []
                for i, sheet in enumerate(book.worksheets):
                    tokens = []
                    for image in sheet._images:
                        with io.BytesIO(image._data()) as stream:
                            tokens.append(vision.writer.token(vision.analyze(stream)))
                    if tokens:
                        end = starts[i + 1] if i + 1 < len(starts) else len(result.markdown)
                        additions.append((end, "\n\n" + "\n\n".join(tokens) + "\n\n"))
                for end, text in reversed(additions):
                    result.markdown = result.markdown[:end] + text + result.markdown[end:]
                return result
            finally:
                book.close()

    return {"docx": Docx, "pptx": Pptx, "xlsx": Xlsx}[fmt]()
