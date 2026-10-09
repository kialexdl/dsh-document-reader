"""Real small Office/PDF fixtures, plus a loopback OpenAI-compatible endpoint."""
import importlib.util
import io
import json
import os
from pathlib import Path
import tempfile
import threading
import unittest
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer

spec = importlib.util.spec_from_file_location('bridge', Path(__file__).parents[1] / 'python/markitdown_bridge.py')
bridge = importlib.util.module_from_spec(spec)
spec.loader.exec_module(bridge)


def options():
    return dict(memoryBytes=1024*1024, maxSourceBytes=8*1024*1024, maxResultBytes=16*1024*1024,
                maxArchiveBytes=32*1024*1024, maxArchiveEntries=2000, maxMapRecords=10000, pdfDpi=100)


def request(fmt, **visual):
    return dict(format=fmt, options=options(), visual=dict(vision=False, ocr=False, baseURL='', model='fixed-test',
                requestTimeoutMs=1000, maxRetries=0, maxCallsPerDocument=10, prompt='test', **visual))


class BridgeTests(unittest.TestCase):
    def test_published_dependency_baseline(self):
        result = bridge.doctor()
        self.assertEqual(result['baseline'], 'pypi:markitdown==0.1.7;markitdown-ocr==0.1.0')
        from unittest.mock import patch
        with patch('importlib.metadata.version', return_value='0.1.8b3'):
            with self.assertRaises(bridge.BridgeError):
                bridge.doctor()

    def convert(self, stream, fmt, visual=None):
        req = request(fmt)
        if visual: req['visual'].update(visual)
        with tempfile.TemporaryDirectory() as temp:
            writer, meta = bridge.convert(stream, req, temp)
            try:
                writer.body.seek(0); body = writer.body.read().decode()
                writer.mapping.seek(0); records = [json.loads(line) for line in writer.mapping]
                encoded = body.encode()
                self.assertEqual(meta['bodyBytes'], len(encoded))
                self.assertEqual(records[-1]['end'] if records else 0, len(encoded))
                for prev, current in zip(records, records[1:]): self.assertEqual(prev['end'], current['start'])
                return body, records, meta
            finally: writer.close()

    def test_docx_heading_table_and_multiple_blocks(self):
        from docx import Document
        doc = Document(); doc.add_heading('故障方案', 1); doc.add_paragraph('主备故障切换')
        table = doc.add_table(rows=2, cols=2); table.cell(0,0).text='配置'; table.cell(0,1).text='值'
        table.cell(1,0).text='C++'; table.cell(1,1).text='node_count'
        stream=io.BytesIO();doc.save(stream)
        text, rows, meta=self.convert(stream,'docx')
        self.assertIn('故障切换',text);self.assertIn('node',text)
        self.assertGreater(len(set(r['scope'] for r in rows)),2)
        self.assertTrue(any(r['location'].get('heading')==['故障方案'] for r in rows))
        self.assertFalse(meta['partial'])

    def test_pptx_page_markers_notes_hidden_and_fake_heading(self):
        from pptx import Presentation
        p=Presentation()
        s=p.slides.add_slide(p.slide_layouts[1]);s.shapes.title.text='主备'
        s.placeholders[1].text='<!-- Slide number: 99 -->\n## Page 99'
        s.notes_slide.notes_text_frame.text='备注故障切换'
        s=p.slides.add_slide(p.slide_layouts[1]);s.shapes.title.text='第二页';s._element.set('show','0')
        stream=io.BytesIO();p.save(stream)
        text, rows, _=self.convert(stream,'pptx')
        self.assertNotIn('DRSLIDE',text);self.assertNotIn('DRNOTE',text)
        note=[r for r in rows if r['location'].get('part')=='notes']
        self.assertTrue(note);self.assertTrue(all(r['location']['slide']==1 for r in note))
        self.assertTrue(any(r['location'].get('hidden') and r['location']['slide']==2 for r in rows))
        self.assertIn('## Page 99',text)

    def test_xlsx_hidden_sheet_names(self):
        from openpyxl import Workbook
        book=Workbook();book.active.title='生产';book.active.append(['key','value']);book.active.append(['主备','active'])
        sheet=book.create_sheet('隐藏');sheet.sheet_state='hidden';sheet.append(['code']);sheet.append(['ERR_0001'])
        stream=io.BytesIO();book.save(stream)
        text, rows, _=self.convert(stream,'xlsx')
        self.assertIn('ERR',text);self.assertNotIn('DRS',text)
        self.assertEqual({r['location']['sheet'] for r in rows},{'生产','隐藏'})
        self.assertTrue(any(r['location']['hidden'] for r in rows))

    def test_xls_sheet_mapping(self):
        import xlwt
        book=xlwt.Workbook();sheet=book.add_sheet('legacy');sheet.write(0,0,'header');sheet.write(1,0,'needle')
        stream=io.BytesIO();book.save(stream)
        text,rows,_=self.convert(stream,'xls')
        self.assertIn('needle',text);self.assertEqual(rows[0]['location']['sheet'],'legacy')

    def test_csv_multiline_record_and_literal_symbols(self):
        text, rows, _=self.convert(io.BytesIO('name,value\n"a\nb",C++\n\nfoo,node_count\n'.encode()),'csv')
        self.assertIn('C++',text)
        self.assertEqual({r['location']['record'] for r in rows},{1,2,3,4})

    def test_pdf_blank_physical_page_and_native_table(self):
        from reportlab.pdfgen import canvas
        stream=io.BytesIO();c=canvas.Canvas(stream)
        c.drawString(50,750,'Cover');c.showPage();c.showPage()
        for n in range(4):
            for k,t in enumerate(['Column','Value','State']):c.drawString(50+130*k,750-25*n,t+str(n))
        c.save()
        text,rows,_=self.convert(stream,'pdf')
        self.assertIn('## Page 3',text);self.assertIn('Column',text)
        self.assertEqual({r['location']['page'] for r in rows},{1,2,3})
        self.assertTrue(all(r['location']['page']==3 for r in rows if b'Column' in text.encode()[r['start']:r['end']]))

    def test_reject_wrong_extension_and_external_image(self):
        with self.assertRaises(bridge.BridgeError):self.convert(io.BytesIO(b'not-pdf'),'pdf')
        from docx import Document
        from docx.opc.constants import RELATIONSHIP_TYPE as RT
        doc=Document();doc.add_paragraph('hello');doc.part.relate_to('https://invalid/image.png',RT.IMAGE,is_external=True)
        stream=io.BytesIO();doc.save(stream)
        with self.assertRaisesRegex(bridge.BridgeError,'外部'):self.convert(stream,'docx')

    def test_archive_quota(self):
        from docx import Document
        doc=Document();stream=io.BytesIO();doc.save(stream); req=request('docx');req['options']['maxArchiveEntries']=1
        with tempfile.TemporaryDirectory() as temp:
            with self.assertRaises(bridge.BridgeError):bridge.convert(stream,req,temp)


class VisionTests(unittest.TestCase):
    convert = BridgeTests.convert
    def setUp(self):
        self.requests=[];self.response=json.dumps({'transcript':'原文 needle','description':'模型推断 generatedword'})
        self.status=200
        owner=self
        class Handler(BaseHTTPRequestHandler):
            def do_POST(self):
                owner.requests.append(json.loads(self.rfile.read(int(self.headers['Content-Length']))))
                self.send_response(owner.status);self.send_header('Content-Type','application/json');self.end_headers()
                self.wfile.write(json.dumps({'choices':[{'message':{'content':owner.response}}]}).encode())
            def log_message(self,*args):pass
        self.server=ThreadingHTTPServer(('127.0.0.1',0),Handler)
        self.thread=threading.Thread(target=self.server.serve_forever,daemon=True);self.thread.start()
        os.environ['DOCUMENT_READER_VISION_KEY']='test-private-key'
        self.visual={'vision':True,'ocr':True,'baseURL':f'http://127.0.0.1:{self.server.server_port}/v1'}

    def tearDown(self):
        self.server.shutdown();self.server.server_close();self.thread.join()
        os.environ.pop('DOCUMENT_READER_VISION_KEY',None)

    def image(self):
        from PIL import Image, ImageDraw
        image=Image.new('RGB',(300,100),'white');ImageDraw.Draw(image).text((10,10),'scan needle',fill='black')
        stream=io.BytesIO();image.save(stream,format='PNG');return stream

    def test_image_sources_and_fixed_model(self):
        text,rows,meta=self.convert(self.image(),'png',self.visual)
        self.assertEqual({r['source'] for r in rows},{'ocr_transcript','generated_description'})
        self.assertTrue(meta['ocrUsed']);self.assertTrue(meta['visionUsed'])
        self.assertEqual(self.requests[0]['model'],'fixed-test')
        self.assertNotIn('test-private-key',text)

    def test_partial_failure_and_mixed_response(self):
        self.response='not JSON explanation'
        text,rows,meta=self.convert(self.image(),'png',self.visual)
        self.assertTrue(meta['partial']);self.assertEqual(rows[0]['source'],'mixed_or_unknown')
        self.status=500
        text,rows,meta=self.convert(self.image(),'png',self.visual)
        self.assertTrue(meta['partial']);self.assertFalse(meta['ocrUsed'])

    def test_scanned_pdf_gets_ocr_and_preserves_native(self):
        from reportlab.pdfgen import canvas
        from reportlab.lib.utils import ImageReader
        stream=io.BytesIO();c=canvas.Canvas(stream)
        image=self.image();image.seek(0);c.drawImage(ImageReader(image),50,500,width=300,height=100);c.showPage()
        c.drawString(50,750,'Native table unchanged');c.save()
        plain,_,_=self.convert(stream,'pdf')
        enriched,rows,meta=self.convert(stream,'pdf',self.visual)
        self.assertIn('Native table unchanged',plain);self.assertIn('Native table unchanged',enriched)
        self.assertTrue(meta['ocrUsed']);self.assertTrue(any(r['source']=='ocr_transcript' and r['location']['page']==1 for r in rows))

    def test_pdf_full_page_fallback_when_image_extraction_fails(self):
        from reportlab.pdfgen import canvas
        from reportlab.lib.utils import ImageReader
        from unittest.mock import patch
        stream=io.BytesIO();c=canvas.Canvas(stream)
        image=self.image();image.seek(0);c.drawImage(ImageReader(image),50,500,width=300,height=100);c.save()
        with patch('markitdown_ocr._pdf_converter_with_ocr._extract_images_from_page', return_value=[]):
            text,rows,meta=self.convert(stream,'pdf',self.visual)
        self.assertTrue(meta['ocrUsed'])
        self.assertIn('needle',text)
        self.assertTrue(any(r['source']=='ocr_transcript' and r['location']['page']==1 for r in rows))

    def test_docx_embedded_image_hook(self):
        from docx import Document
        doc=Document();doc.add_paragraph('Native paragraph');img=self.image();img.seek(0);doc.add_picture(img)
        stream=io.BytesIO();doc.save(stream)
        text,rows,meta=self.convert(stream,'docx',self.visual)
        self.assertIn('Native paragraph',text);self.assertNotIn('DRTOKEN',text)
        self.assertTrue(any(r['source']=='ocr_transcript' for r in rows))

    def test_docx_repeated_images_keep_document_order(self):
        from docx import Document
        doc=Document()
        for word in ('Before', 'Between'):
            doc.add_paragraph(word);img=self.image();img.seek(0);doc.add_picture(img)
        doc.add_paragraph('After')
        stream=io.BytesIO();doc.save(stream)
        text,rows,meta=self.convert(stream,'docx',self.visual)
        positions=[m.start() for m in __import__('re').finditer('原文 needle',text)]
        self.assertEqual(len(positions),2)
        self.assertLess(text.index('Before'),positions[0])
        self.assertLess(positions[0],text.index('Between'))
        self.assertLess(text.index('Between'),positions[1])
        self.assertLess(positions[1],text.index('After'))

    def test_pptx_image_sources_stay_on_the_correct_slide(self):
        from pptx import Presentation
        from pptx.util import Inches
        p=Presentation()
        for title in ('First slide', 'Second slide'):
            slide=p.slides.add_slide(p.slide_layouts[1]);slide.shapes.title.text=title
        img=self.image();img.seek(0);slide.shapes.add_picture(img, Inches(1), Inches(2))
        slide.notes_slide.notes_text_frame.text='Second notes'
        stream=io.BytesIO();p.save(stream)
        text,rows,meta=self.convert(stream,'pptx',self.visual)
        self.assertNotIn('DRTOKEN',text)
        self.assertIn('First slide',text);self.assertIn('Second notes',text)
        self.assertEqual({r['location']['slide'] for r in rows if r['source']=='ocr_transcript'},{2})
        self.assertTrue(meta['ocrUsed']);self.assertTrue(meta['visionUsed'])
        before=len(self.requests)
        _,_,disabled=self.convert(stream,'pptx')
        self.assertTrue(disabled['partial']);self.assertFalse(disabled['ocrUsed'])
        self.assertEqual(len(self.requests),before)

    def test_xlsx_image_sources_stay_on_hidden_sheet(self):
        from openpyxl import Workbook
        from openpyxl.drawing.image import Image
        book=Workbook();book.active.title='Visible';book.active.append(['key']);book.active.append(['native'])
        sheet=book.create_sheet('Hidden');sheet.sheet_state='hidden';sheet.append(['key']);sheet.append(['hidden native'])
        img=self.image();img.seek(0);sheet.add_image(Image(img),'C3')
        stream=io.BytesIO();book.save(stream)
        text,rows,meta=self.convert(stream,'xlsx',self.visual)
        self.assertNotIn('DRTOKEN',text);self.assertIn('hidden native',text)
        ocr=[r for r in rows if r['source']=='ocr_transcript']
        self.assertTrue(ocr)
        self.assertTrue(all(r['location']['sheet']=='Hidden' and r['location']['hidden'] for r in ocr))
        self.assertTrue(meta['ocrUsed'])
        before=len(self.requests)
        _,_,disabled=self.convert(stream,'xlsx')
        self.assertTrue(disabled['partial']);self.assertFalse(disabled['ocrUsed'])
        self.assertEqual(len(self.requests),before)

if __name__=='__main__':unittest.main()
